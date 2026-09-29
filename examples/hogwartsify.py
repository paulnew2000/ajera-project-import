"""
Turn an export of real projects into a safe, obviously-fake test workbook (Harry Potter theme).

    python examples/hogwartsify.py export.xlsx test.xlsx
        Replaces Job IDs (HP001...), descriptions, notes, dates, managers/principals and phase
        names. KEEPS Client / Department / Project Type / Rate Table / Invoice Format so the file
        still validates against YOUR Ajera - use this one to test creating projects.

    python examples/hogwartsify.py export.xlsx sample.xlsx --scrub-references
        Also replaces every Ajera name (clients, departments, ...) and the Lookups sheet with
        made-up ones. Safe to share, but it will NOT validate against a real Ajera until you
        change those columns to your own names.

The script prints counts only - never the original values.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from workbook import build_workbook  # noqa: E402

CHARACTERS = ["Dumbledore", "McGonagall", "Hagrid", "Hermione Granger", "Snape", "Luna Lovegood",
              "Neville Longbottom", "Sirius Black", "Remus Lupin", "Nymphadora Tonks", "Ginny Weasley",
              "Dobby", "Minerva's Cat", "Mad-Eye Moody", "Arthur Weasley", "Fleur Delacour"]
NOTES = [
    "Test import example - owl post to the client on project start.",
    "Test import example - Fawkes to confirm scope before field work.",
    "Test import example - keep away from the Whomping Willow during site visits.",
    "Test import example - invoice in Galleons, Sickles and Knuts.",
    "Test import example - Room of Requirement booked for kickoff meeting.",
    "Test import example - Ministry permit required before any apparition on site.",
    "Test import example - Hufflepuff team leads sampling; badgers welcome.",
    "Test import example - Marauder's Map attached to the proposal.",
    "Test import example - no Floo travel billed without approval.",
    "Test import example - Sorting Hat assigns staff at kickoff.",
]
PHASES = ["Task 1 - Owl Post Kickoff", "Task 2 - Potions Lab Analysis", "Task 3 - Forbidden Forest Field Work",
          "Task 4 - Charms Report", "Task 5 - Quidditch Pitch Survey", "Task 6 - Great Hall Presentation",
          "Task 7 - Hogsmeade Site Visit", "Task 8 - Room of Requirement Review"]
# Fake reference names for --scrub-references (each distinct real value gets the next name)
FAKE = {
    "Client": ["Gringotts Wizarding Bank", "Ministry of Magic", "Hogwarts School of Witchcraft and Wizardry",
               "Ollivanders Wand Shop", "Weasleys' Wizard Wheezes", "Honeydukes Sweetshop", "The Leaky Cauldron",
               "St Mungo's Hospital", "The Daily Prophet", "Flourish and Blotts", "Borgin and Burkes", "Madam Malkin's"],
    "Department": ["Gryffindor", "Hufflepuff", "Ravenclaw", "Slytherin", "Hogsmeade", "Diagon Alley"],
    "Project Type": ["Potions", "Charms", "Transfiguration", "Herbology", "Care of Magical Creatures", "Divination",
                     "Astronomy", "Defence Against the Dark Arts"],
    "Rate Table": ["Standard Rates 2026", "Ministry Rates 2026", "Standard Rates 2025", "Goblin Rates 2026",
                   "Hogwarts Staff Rates", "Order of the Phoenix Rates"],
    "Invoice Format": ["Standard Invoice", "Detailed Invoice", "Summary Invoice", "Ministry Invoice",
                       "Goblin Invoice", "Owl-Post Invoice", "Howler Invoice", "Parchment Invoice"],
}
FAKE_PEOPLE = {"Project Manager": ["Albus Dumbledore", "Minerva McGonagall", "Filius Flitwick", "Pomona Sprout"],
               "Principal in Charge": ["Albus Dumbledore", "Minerva McGonagall"]}
FIXED_FEE_ROWS = {2, 5, 8}  # vary billing types so the example shows more than one


def _rows(ws):
    heads = [c.value for c in ws[1]]
    return [dict(zip(heads, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("dest")
    ap.add_argument("--scrub-references", action="store_true")
    ap.add_argument("--prefix", default="HP")
    ap.add_argument("--start", type=lambda s: date.fromisoformat(s), default=date(2026, 11, 2))
    a = ap.parse_args()

    wb = load_workbook(a.source, data_only=True)
    projects, phases = _rows(wb["Projects"]), _rows(wb["Phases"])
    if len(projects) > len(CHARACTERS):
        sys.exit(f"Only {len(CHARACTERS)} character names available; the file has {len(projects)} projects.")

    maps: dict[str, dict] = {k: {} for k in FAKE}

    def fake(col, value):
        if not a.scrub_references or value in (None, ""):
            return value
        m = maps[col]
        if value not in m:
            m[value] = FAKE[col][len(m) % len(FAKE[col])]
        return m[value]

    new_id = {}
    for i, p in enumerate(projects):
        old = p["Job ID"]
        new_id[old] = f"{a.prefix}{i + 1:03d}"
        start = a.start + timedelta(weeks=2 * i)
        p.update({
            "Job ID": new_id[old],
            "Description": f"Test Project: {CHARACTERS[i]}",
            "Billing Type": "Fixed Fee" if i in FIXED_FEE_ROWS else "Time and Expense",
            "Project Manager": None, "Principal in Charge": None,
            "Status": "Active",
            "Start Date": start, "End Date": start + timedelta(days=90 + 30 * (i % 4)),
            "Invoice Group": "Invoice Group 1",
            "Notes": NOTES[i % len(NOTES)],
        })
        for col in FAKE:
            p[col] = fake(col, p.get(col))

    out_phases, counter = [], {}
    for ph in phases:
        jid = new_id.get(ph["Job ID"])
        if jid is None:
            continue
        n = counter[jid] = counter.get(jid, 0) + 1
        out_phases.append({"Job ID": jid, "Phase Description": PHASES[(n - 1) % len(PHASES)],
                           "Department": fake("Department", ph.get("Department")),
                           "Billing Type": None})  # inherit the project's billing type

    if a.scrub_references:
        # A made-up Lookups sheet consistent with the fake names used above
        lookups = {col: [(n, "") for n in FAKE[col]] for col in FAKE}
        lookups["managers"] = [(n, "") for n in FAKE_PEOPLE["Project Manager"]]
        lookups["principals"] = [(n, "") for n in FAKE_PEOPLE["Principal in Charge"]]
        note = "Sample data (Harry Potter theme). Replace Client/Department/Type/Rate Table with names from YOUR Ajera."
    else:
        lk_ws = wb["Lookups"]
        heads = [c.value for c in lk_ws[1]]
        cols = {h: i for i, h in enumerate(heads) if h and h != "Key"}
        src = {"Client": "Client", "Department": "Department", "Project Type": "Project Type",
               "Rate Table": "Rate Table", "Invoice Format": "Invoice Format",
               "managers": "Project Manager", "principals": "Principal in Charge"}
        lookups = {k: [(r[cols[h]], r[cols[h] + 1]) for r in lk_ws.iter_rows(min_row=2, values_only=True)
                       if r[cols[h]] not in (None, "")] for k, h in src.items() if h in cols}
        note = "Test copy (Harry Potter theme) - descriptions, IDs, dates and phases are fake; Ajera names are real."

    build_workbook(Path(a.dest), projects, out_phases, lookups, "Ajera Project Import - Hogwarts test data", note)
    print(f"Wrote {a.dest}: {len(projects)} projects, {len(out_phases)} phases, "
          f"references {'scrubbed' if a.scrub_references else 'kept'}"
          + (f" ({', '.join(f'{k} {len(v)}' for k, v in maps.items())} distinct values replaced)" if a.scrub_references else ""))


if __name__ == "__main__":
    main()
