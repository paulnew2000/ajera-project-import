"""
Make the Hogwarts sample runnable in YOUR Ajera.

The sample's clients, departments, project types and rate tables are made up (Gringotts,
Gryffindor, ...), so they don't exist in your database. This rewrites those columns with
valid values from your Ajera and keeps the Harry Potter descriptions, notes, dates and phases:

    Client          one client you choose (a test/placeholder client is best)
    Department      each sample department -> one of your departments (spread across them)
    Project Type    each sample type -> one of your project types (overhead/marketing types avoided)
    Rate Table      one rate table (you choose, or the newest one with "standard" in its name)
    Invoice Format  left blank (optional column)
    Job ID          HP001... continuing after any HP numbers already in your Ajera
    Lookups sheet   replaced with your Ajera's lists, so the dropdowns work

Nothing is written to Ajera - it only reads lists and saves a new workbook.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

import importer
import workbook
from importer import Lookups, job_id_text

HERE = Path(__file__).resolve().parent
SAMPLE = HERE / "examples" / "Hogwarts sample projects.xlsx"
DEFAULT_OUT = HERE / "Hogwarts sample - ready for my Ajera.xlsx"
TEST_CLIENT = re.compile(r"test|sample|demo|placeholder|new client|do not use", re.I)
AVOID_TYPE = re.compile(r"overhead|marketing|admin|internal|proposal|leave|holiday", re.I)


def _rows(ws) -> list[dict]:
    heads = [c.value for c in ws[1]]
    return [dict(zip(heads, r)) for r in ws.iter_rows(min_row=2, values_only=True)
            if any(v not in (None, "") for v in r)]


def _pick_client(lk: Lookups, choice: str) -> tuple[int, str]:
    if choice:
        key, why = lk.find("Client", choice)
        if why:
            raise ValueError(f"Client you asked for: {why}")
        return key, "the client you chose"
    hits = [(n, k) for n, k in lk.rows["Client"] if TEST_CLIENT.search(n)]
    if len(hits) == 1:
        return hits[0][1], "auto-picked: the only client whose name looks like a test/placeholder client"
    raise ValueError(
        ("Several" if hits else "No") + " clients look like a test/placeholder client. Say which client the sample "
        "projects should use: its name or Key from the Lookups sheet of a template (menu option 2).")


def _pick_rate_table(lk: Lookups, choice: str) -> tuple[int, str]:
    if choice:
        key, why = lk.find("Rate Table", choice)
        if why:
            raise ValueError(f"Rate table you asked for: {why}")
        return key, "the rate table you chose"
    tables = lk.rows["Rate Table"]
    if not tables:
        raise ValueError("No active rate tables found in Ajera.")
    std = [(n, k) for n, k in tables if "standard" in n.lower()]
    if std:
        def year(n):
            ys = re.findall(r"\b(20\d\d)\b", n)
            return max(map(int, ys)) if ys else 0
        # newest year first; then the plainest (shortest) name - "2026 Standard" beats "Acme 2026 Standard Rates"
        _, k = min(std, key=lambda t: (-year(t[0]), len(t[0]), t[1]))
        return k, "auto-picked: newest, plainest rate table with 'standard' in its name"
    return tables[0][1], "auto-picked: first active rate table (none is named 'standard')"


def _spread(sample_values: list, targets: list[tuple[str, int]]) -> dict:
    """Map each distinct sample value to a target name, cycling through the targets."""
    out = {}
    for v in sample_values:
        if v not in (None, "") and v not in out:
            out[v] = targets[len(out) % len(targets)][0]
    return out


def _next_number(existing_ids: set[str], prefix: str) -> int:
    pat = re.compile(re.escape(prefix.lower()) + r"(\d+)")
    nums = [int(m.group(1)) for i in existing_ids if (m := pat.fullmatch(i))]
    return max(nums, default=0) + 1


def adapt_sample(aj, client: str = "", rate_table: str = "", prefix: str = "HP",
                 source: Path = SAMPLE, dest: Path = DEFAULT_OUT) -> tuple[Path, list[str], list]:
    """Returns (workbook path, summary lines, CHECK problems)."""
    lk = importer.load_lookups(aj)
    client_key, client_why = _pick_client(lk, client)
    rate_key, rate_why = _pick_rate_table(lk, rate_table)
    types = [t for t in lk.rows["Project Type"] if not AVOID_TYPE.search(t[0])] or lk.rows["Project Type"]
    if not lk.rows["Department"] or not types:
        raise ValueError("Your Ajera returned no active departments or project types.")

    wb = load_workbook(source, data_only=True)
    projects, phases = _rows(wb["Projects"]), _rows(wb["Phases"])
    dept_map = _spread([p.get("Department") for p in projects] + [p.get("Department") for p in phases],
                       lk.rows["Department"])
    type_map = _spread([p.get("Project Type") for p in projects], types)
    start = _next_number(lk.existing_ids, prefix)
    new_id = {}
    for i, p in enumerate(projects):
        new_id[job_id_text(p["Job ID"])] = f"{prefix}{start + i:03d}"
        p.update({
            "Job ID": new_id[job_id_text(p["Job ID"])],
            "Client": lk.name_of("Client", client_key),
            "Department": dept_map.get(p.get("Department")),
            "Project Type": type_map.get(p.get("Project Type")),
            "Rate Table": lk.name_of("Rate Table", rate_key),
            "Invoice Format": None,
            "Project Manager": None, "Principal in Charge": None,
        })
    for ph in phases:
        ph["Job ID"] = new_id.get(job_id_text(ph["Job ID"]), ph["Job ID"])
        if ph.get("Department"):
            ph["Department"] = dept_map.get(ph["Department"])

    workbook.build_workbook(Path(dest), projects, phases, workbook.lookup_rows_from(lk),
                            "Ajera Project Import - Hogwarts sample (ready for your Ajera)",
                            "Harry Potter descriptions and phases; Client/Department/Type/Rate Table filled from your Ajera.")

    rows, problems = importer.read_workbook(dest)
    if not problems:
        problems = importer.validate(rows, lk)
    summary = [
        f"Client: Ajera key {client_key} ({client_why})",
        f"Rate table: Ajera key {rate_key} ({rate_why})",
        f"Departments: {len(dept_map)} sample departments spread over {min(len(dept_map), len(lk.rows['Department']))} of yours",
        f"Project types: {len(type_map)} sample types spread over {min(len(type_map), len(types))} of yours",
        f"Job IDs: {prefix}{start:03d} to {prefix}{start + len(projects) - 1:03d} "
        f"({len(projects)} projects, {len(phases)} phases)",
        "Invoice Format, Project Manager and Principal left blank (optional).",
    ]
    return Path(dest), summary, problems
