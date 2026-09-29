"""Offline tests: synthetic lookups + the sample workbook. No Ajera connection needed.

    python -m pytest -q
"""
from pathlib import Path

import pytest
from openpyxl import load_workbook

import importer
from importer import Lookups, read_workbook, validate

SAMPLE = Path(__file__).parent / "examples" / "Hogwarts sample projects.xlsx"


def fake_lookups(existing=()):
    wb = load_workbook(SAMPLE, data_only=True)
    ws = wb["Lookups"]
    heads = [c.value for c in ws[1]]
    rows = {}
    key = 100
    for i, h in enumerate(heads):
        if h in ("Client", "Department", "Project Type", "Rate Table", "Invoice Format"):
            names = [r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i]]
            rows[h] = [(n, key + j) for j, n in enumerate(names)]
            key += 100
    rows["Employee"] = [("Albus Dumbledore", 1), ("Minerva McGonagall", 2), ("Argus Filch", 3)]
    return Lookups(company_key=1, rows=rows, managers={1, 2}, principals={1},
                   existing_ids={e.lower() for e in existing})


def test_sample_validates_against_its_own_lookups():
    projects, problems = read_workbook(SAMPLE)
    assert not problems
    assert len(projects) == 10
    assert validate(projects, fake_lookups()) == []
    p = projects[0].args
    assert p["Project"]["ID"] == "HP001"
    assert p["Project"]["Description"].startswith("Test Project:")
    assert p["Project"]["RateTableKey"] and p["InvoiceGroups"][0]["ClientKey"]
    assert "ProjectManager" not in p["Project"]  # blank = left out, not sent as null
    assert all(ph["Description"] and ph["DepartmentKey"] for ph in p["Phases"])


def test_existing_and_duplicate_ids_are_rejected():
    projects, _ = read_workbook(SAMPLE)
    projects[1].values["Job ID"] = "HP001"
    probs = [str(p) for p in validate(projects, fake_lookups(existing=["HP003"]))]
    assert any("row 3, Job ID: is also used on row 2" in p for p in probs)
    assert any("row 4, Job ID: already exists in Ajera" in p for p in probs)


def test_problem_messages_never_echo_cell_values():
    projects, _ = read_workbook(SAMPLE)
    projects[0].values["Client"] = "Secret Client LLC"
    projects[0].values["Project Manager"] = "Argus Filch"  # exists but is not a PM
    probs = [str(p) for p in validate(projects, fake_lookups())]
    assert probs and not any("Secret" in p or "Filch" in p for p in probs)
    assert any("not flagged as a Project Manager" in p for p in probs)


@pytest.mark.parametrize("value,expected", [("Fixed Fee", "FixedFee"), ("fixedfee", "FixedFee"),
                                            ("Time and Expense", "TimeAndExpense"), ("bogus", None)])
def test_billing_type_labels(value, expected):
    assert importer._billing(value) == expected


def test_project_without_phases_gets_a_default_phase():
    projects, _ = read_workbook(SAMPLE)
    projects[0].phases = []
    validate(projects, fake_lookups())
    assert projects[0].args["Phases"][0]["Description"] == importer.DEFAULT_PHASE


def test_number_job_ids_match_phases():
    assert importer.job_id_text(11001.0) == "11001"
    assert importer.job_id_text(" HP001 ") == "HP001"
