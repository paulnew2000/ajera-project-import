"""
Excel -> Ajera project importer: the shared engine behind the command line (cli.py) and the
MCP server (server.py).

    read workbook -> look up names in Ajera -> check every row -> (only if all rows pass)
    create each project -> read it back -> log it -> write a results copy of the workbook

Messages never repeat cell values (they point at sheet/row/column instead), so output can be
pasted into a ticket or chat without leaking client or employee names.
"""
from __future__ import annotations

import csv
import os
import re
import shutil
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ajera_client import Ajera, AjeraError

HERE = Path(__file__).resolve().parent

# -- workbook layout ---------------------------------------------------------------------
PROJECT_COLUMNS = [
    # (header, required, help text shown on the Instructions sheet)
    ("Job ID", True, "Your project number. Must not already exist in Ajera."),
    ("Description", True, "Project name as it should appear in Ajera."),
    ("Client", True, "Client name exactly as in Ajera (see Lookups sheet)."),
    ("Department", True, "Department name (see Lookups sheet)."),
    ("Project Type", True, "Project type (see Lookups sheet)."),
    ("Billing Type", True, "Time and Expense, Fixed Fee, Percent Complete, Unit Price, Nonbillable, ..."),
    ("Rate Table", True, "Billing rate table name (see Lookups sheet). Ajera requires one."),
    ("Project Manager", False, "Optional. Employee name; must be flagged as a Project Manager in Ajera."),
    ("Principal in Charge", False, "Optional. Employee name; must be flagged as a Principal in Ajera."),
    ("Status", False, "Optional. Active (default) or Preliminary."),
    ("Start Date", False, "Optional. Estimated start date."),
    ("End Date", False, "Optional. Estimated completion date."),
    ("Invoice Group", False, "Optional. Name of the invoice group (default: Invoice Group 1)."),
    ("Invoice Format", False, "Optional. Invoice format name (see Lookups sheet)."),
    ("Notes", False, "Optional. Free text copied into the project's Notes."),
]
PHASE_COLUMNS = [
    ("Job ID", True, "Which project this phase belongs to (matches the Projects sheet)."),
    ("Phase Description", True, "Phase name, e.g. Task 1 or Field Work."),
    ("Department", False, "Optional. Defaults to the project's department."),
    ("Billing Type", False, "Optional. Defaults to the project's billing type."),
]

BILLING_TYPES = {  # friendly label -> Ajera value
    "Time and Expense": "TimeAndExpense",
    "Fixed Fee": "FixedFee",
    "Percent Complete": "PercentComplete",
    "Unit Price": "UnitPrice",
    "Percent of Construction Cost": "PercentofConstructionCost",
    "Nonbillable": "Nonbillable",
    "Marketing": "Marketing",
    "Overhead": "Overhead",
}
BILLING_LABEL = {v: k for k, v in BILLING_TYPES.items()}
STATUSES = ["Active", "Preliminary"]
DEFAULT_PHASE = "General"
DEFAULT_INVOICE_GROUP = "Invoice Group 1"


def created_log_path() -> Path:
    p = Path(os.environ.get("AJERA_CREATED_LOG") or "created_projects.csv")
    return p if p.is_absolute() else HERE / p


# -- lookups -----------------------------------------------------------------------------
def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _first_list(content: dict) -> list:
    return next((v for v in content.values() if isinstance(v, list)), [])


def _person_name(e: dict) -> str:
    return " ".join(p for p in (e.get("FirstName"), e.get("MiddleName"), e.get("LastName")) if p).strip()


def _desc(r: dict) -> str:
    return r.get("Description") or r.get("Name") or r.get("Department") or ""


def _key(r: dict, *names: str):
    for n in names:
        if r.get(n) is not None:
            return r[n]
    return next((v for k, v in r.items() if k.endswith("Key") and isinstance(v, int)), None)


@dataclass
class Lookups:
    """Name -> key tables read from Ajera. `rows` keeps (name, key) lists for the Lookups sheet."""
    company_key: int
    rows: dict[str, list[tuple[str, int]]]
    managers: set[int]
    principals: set[int]
    existing_ids: set[str]

    def find(self, kind: str, value) -> tuple[int | None, str]:
        """Return (key, problem). A number in the cell is taken as the Ajera key itself."""
        table = self.rows[kind]
        if isinstance(value, (int, float)) or re.fullmatch(r"\d+", str(value).strip()):
            k = int(value)
            return (k, "") if any(key == k for _, key in table) else (None, "key number not found in Ajera")
        hits = [key for name, key in table if _norm(name) == _norm(value)]
        if not hits:
            return None, "not found in Ajera (check spelling against the Lookups sheet)"
        if len(set(hits)) > 1:
            return None, "matches more than one Ajera record - type the key number from the Lookups sheet instead"
        return hits[0], ""

    def name_of(self, kind: str, key) -> str:
        return next((n for n, k in self.rows[kind] if k == key), str(key) if key is not None else "")


def load_lookups(aj: Ajera, want_existing_ids: bool = True) -> Lookups:
    companies = aj.call("ListCompanies").get("Companies", [])
    depts = [(d.get("Department") or _desc(d), d["DepartmentKey"])
             for d in aj.call("ListDepartments").get("Departments", []) if d.get("Status", "Active") == "Active"]
    ptypes = [(_desc(t), t["ProjectTypeKey"])
              for t in aj.call("ListProjectTypes", {"FilterByStatus": ["Active"]}).get("ProjectTypes", [])]
    clients = [(_desc(c), c["ClientKey"])
               for c in aj.call("ListClients", {"FilterByStatus": ["Active"]}).get("Clients", [])]
    rates = [(_desc(r), r["RateTableKey"])
             for r in aj.call("ListRateTables").get("RateTables", []) if r.get("Status", "Active") == "Active"]
    try:
        formats = [(_desc(f), _key(f, "InvoiceFormatKey"))
                   for f in _first_list(aj.call("ListInvoiceFormats")) if f.get("Status", "Active") == "Active"]
    except AjeraError:
        formats = []  # method not enabled for this API user: the column just can't be used
    emps = aj.call("ListEmployees", {"FilterByStatus": ["Active"]}).get("Employees", [])
    detail = aj.get_employees([e["EmployeeKey"] for e in emps])
    people = [(_person_name(e), e["EmployeeKey"]) for e in detail]
    ids = {str(p.get("ID", "")).strip().lower() for p in aj.list_all_projects()} if want_existing_ids else set()
    return Lookups(
        company_key=companies[0]["CompanyKey"],
        rows={"Client": sorted(clients), "Department": sorted(depts), "Project Type": sorted(ptypes),
              "Rate Table": sorted(rates), "Invoice Format": sorted(formats), "Employee": sorted(people)},
        managers={e["EmployeeKey"] for e in detail if e.get("IsProjectManager")},
        principals={e["EmployeeKey"] for e in detail if e.get("IsPrincipal")},
        existing_ids=ids,
    )


# -- reading the workbook ------------------------------------------------------------------
@dataclass
class Problem:
    sheet: str
    row: int
    column: str
    message: str

    def __str__(self):
        where = f"{self.sheet} row {self.row}" if self.row else self.sheet
        return f"{where}, {self.column}: {self.message}" if self.column else f"{where}: {self.message}"


@dataclass
class ProjectRow:
    row: int
    values: dict
    phases: list[dict] = field(default_factory=list)
    args: dict | None = None  # CreateProjects arguments once validated


def job_id_text(v) -> str:
    """Excel stores 11001 as a number (sometimes 11001.0); Ajera IDs are text."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v if v is not None else "").strip()


def _read_sheet(ws, columns) -> tuple[list[tuple[int, dict]], list[Problem]]:
    headers = [_norm(c.value) for c in ws[1]]
    problems, idx = [], {}
    for name, required, _ in columns:
        if _norm(name) in headers:
            idx[name] = headers.index(_norm(name))
        elif required:
            problems.append(Problem(ws.title, 1, name, "column heading is missing"))
    rows = []
    for r in ws.iter_rows(min_row=2):
        vals = {name: r[i].value if i < len(r) else None for name, i in idx.items()}
        vals = {k: (v.strip() if isinstance(v, str) else v) for k, v in vals.items()}
        if any(v not in (None, "") for v in vals.values()):
            rows.append((r[0].row, vals))
    return rows, problems


def _as_date(v) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError


def _billing(v) -> str | None:
    s = _norm(v).replace(" ", "")
    for label, value in BILLING_TYPES.items():
        if s in (label.lower().replace(" ", ""), value.lower()):
            return value
    return None


def read_workbook(path: str | Path) -> tuple[list[ProjectRow], list[Problem]]:
    wb = load_workbook(path, data_only=True)
    if "Projects" not in wb.sheetnames:
        return [], [Problem("Workbook", 0, "", "there is no sheet named 'Projects'")]
    rows, problems = _read_sheet(wb["Projects"], PROJECT_COLUMNS)
    projects = [ProjectRow(r, v) for r, v in rows]
    if "Phases" in wb.sheetnames:
        prow, pprob = _read_sheet(wb["Phases"], PHASE_COLUMNS)
        problems += pprob
        by_id = {_norm(job_id_text(p.values.get("Job ID"))): p for p in projects}
        for r, v in prow:
            owner = by_id.get(_norm(job_id_text(v.get("Job ID"))))
            if owner is None:
                problems.append(Problem("Phases", r, "Job ID", "does not match any Job ID on the Projects sheet"))
            else:
                owner.phases.append(v | {"_row": r})
    return projects, problems


# -- validation + payload ----------------------------------------------------------------
def validate(projects: list[ProjectRow], lk: Lookups) -> list[Problem]:
    """Check every row and build its CreateProjects arguments. Nothing is sent to Ajera."""
    problems: list[Problem] = []
    seen: dict[str, int] = {}

    def need(pr, sheet, row, col, kind=None):
        v = pr.get(col)
        if v in (None, ""):
            problems.append(Problem(sheet, row, col, "is required"))
            return None
        if kind is None:
            return v
        key, why = lk.find(kind, v)
        if why:
            problems.append(Problem(sheet, row, col, why))
        return key

    for p in projects:
        v, r, before = p.values, p.row, len(problems)
        job_id = job_id_text(need(v, "Projects", r, "Job ID"))
        if job_id:
            if len(job_id) > 30:
                problems.append(Problem("Projects", r, "Job ID", "is longer than 30 characters"))
            if _norm(job_id) in lk.existing_ids:
                problems.append(Problem("Projects", r, "Job ID", "already exists in Ajera"))
            if _norm(job_id) in seen:
                problems.append(Problem("Projects", r, "Job ID", f"is also used on row {seen[_norm(job_id)]}"))
            seen[_norm(job_id)] = r
        desc = need(v, "Projects", r, "Description")
        client = need(v, "Projects", r, "Client", "Client")
        dept = need(v, "Projects", r, "Department", "Department")
        ptype = need(v, "Projects", r, "Project Type", "Project Type")
        rate = need(v, "Projects", r, "Rate Table", "Rate Table")
        billing = None
        if need(v, "Projects", r, "Billing Type") is not None:
            billing = _billing(v["Billing Type"])
            if not billing:
                problems.append(Problem("Projects", r, "Billing Type", "must be one of: " + ", ".join(BILLING_TYPES)))

        pm = pic = fmt = None
        if v.get("Project Manager"):
            pm, why = lk.find("Employee", v["Project Manager"])
            if why:
                problems.append(Problem("Projects", r, "Project Manager", why))
            elif pm not in lk.managers:
                problems.append(Problem("Projects", r, "Project Manager", "this employee is not flagged as a Project Manager in Ajera"))
        if v.get("Principal in Charge"):
            pic, why = lk.find("Employee", v["Principal in Charge"])
            if why:
                problems.append(Problem("Projects", r, "Principal in Charge", why))
            elif pic not in lk.principals:
                problems.append(Problem("Projects", r, "Principal in Charge", "this employee is not flagged as a Principal in Ajera"))
        if v.get("Invoice Format"):
            fmt, why = lk.find("Invoice Format", v["Invoice Format"])
            if why:
                problems.append(Problem("Projects", r, "Invoice Format", why))

        status = (v.get("Status") or "Active")
        status = next((s for s in STATUSES if _norm(s) == _norm(status)), None)
        if not status:
            problems.append(Problem("Projects", r, "Status", "must be Active or Preliminary"))
        dates = {}
        for col in ("Start Date", "End Date"):
            try:
                dates[col] = _as_date(v.get(col))
            except ValueError:
                problems.append(Problem("Projects", r, col, "is not a date (use YYYY-MM-DD or an Excel date)"))
        if dates.get("Start Date") and dates.get("End Date") and dates["End Date"] < dates["Start Date"]:
            problems.append(Problem("Projects", r, "End Date", "is before the Start Date"))

        phases = []
        for ph in p.phases:
            pr = ph["_row"]
            pdesc = need(ph, "Phases", pr, "Phase Description")
            pdept = dept
            if ph.get("Department"):
                pdept, why = lk.find("Department", ph["Department"])
                if why:
                    problems.append(Problem("Phases", pr, "Department", why))
            pbill = billing
            if ph.get("Billing Type"):
                pbill = _billing(ph["Billing Type"])
                if not pbill:
                    problems.append(Problem("Phases", pr, "Billing Type", "must be one of: " + ", ".join(BILLING_TYPES)))
            phases.append({"Description": str(pdesc or ""), "DepartmentKey": pdept, "BillingType": pbill})

        if len(problems) > before:
            continue
        project = {
            "ID": job_id, "Description": str(desc), "CompanyKey": lk.company_key,
            "DepartmentKey": dept, "ProjectTypeKey": ptype, "BillingType": billing,
            "Status": status, "RateTableKey": rate,
        }
        if pm:
            project["ProjectManager"] = {"EmployeeKey": pm}
        if pic:
            project["PrincipalInCharge"] = {"EmployeeKey": pic}
        if dates.get("Start Date"):
            project["EstimatedStartDate"] = dates["Start Date"].isoformat()
        if dates.get("End Date"):
            project["EstimatedCompletionDate"] = dates["End Date"].isoformat()
        if v.get("Notes"):
            project["Notes"] = str(v["Notes"])
        group = {"Description": str(v.get("Invoice Group") or DEFAULT_INVOICE_GROUP), "ClientKey": client}
        if fmt:
            group["InvoiceFormatKey"] = fmt
        p.args = {
            "CreateType": "Project",
            "Project": project,
            "InvoiceGroups": [group],
            "Phases": phases or [{"Description": DEFAULT_PHASE, "DepartmentKey": dept, "BillingType": billing}],
        }
    return problems


# -- creating ------------------------------------------------------------------------------
@dataclass
class Outcome:
    row: int
    job_id: str
    ok: bool
    project_key: int | None = None
    message: str = ""


def _log_created(aj: Ajera, key: int, job_id: str, description: str, source: Path) -> None:
    path = created_log_path()
    new = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["created_at", "database_id", "project_key", "job_id", "description", "source_file"])
        w.writerow([datetime.now().isoformat(timespec="seconds"), aj.database_id, key, job_id, description, source.name])


def create_all(aj: Ajera, projects: list[ProjectRow], source: Path) -> list[Outcome]:
    """Create each validated project, verify it by reading it back, and log it the moment Ajera
    returns a key (the API cannot delete projects, so the log is how you find them again)."""
    aj.check_write_target()
    out = []
    for p in projects:
        job_id = p.args["Project"]["ID"]
        try:
            created = aj.create_project(p.args)
            key = created["ProjectKey"]
            _log_created(aj, key, job_id, p.args["Project"]["Description"], source)
            back = aj.get_projects([key])
            if not back or str(back[0].get("ID")) != job_id:
                out.append(Outcome(p.row, job_id, False, key, "created, but the read-back did not match - check it in Ajera"))
            else:
                out.append(Outcome(p.row, job_id, True, key, "created"))
        except AjeraError as e:
            out.append(Outcome(p.row, job_id, False, None, f"Ajera refused it: {e}"))
        except Exception as e:  # noqa: BLE001 - keep going; report the class only
            out.append(Outcome(p.row, job_id, False, None, f"failed ({type(e).__name__})"))
    return out


def write_results(source: Path, outcomes: list[Outcome]) -> Path:
    """Save a copy of the workbook with Result / Ajera Project Key / Message columns added."""
    dest = source.with_name(f"{source.stem}_results_{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copyfile(source, dest)
    wb = load_workbook(dest)
    ws = wb["Projects"]
    base = ws.max_column + 1
    for i, h in enumerate(("Result", "Ajera Project Key", "Message")):
        ws.cell(1, base + i, h).font = copy(ws.cell(1, 1).font)
    for o in outcomes:
        ws.cell(o.row, base, "Created" if o.ok else "NOT created")
        ws.cell(o.row, base + 1, o.project_key)
        ws.cell(o.row, base + 2, o.message)
    wb.save(dest)
    return dest


# -- one-call entry points (used by cli.py and server.py) ----------------------------------
def check_file(aj: Ajera, path: str | Path) -> tuple[list[ProjectRow], list[Problem], Lookups]:
    projects, problems = read_workbook(path)
    lk = load_lookups(aj)
    if not problems:
        problems += validate(projects, lk)
    if not projects and not problems:
        problems.append(Problem("Projects", 0, "", "no project rows found under the heading row"))
    return projects, problems, lk


def import_file(aj: Ajera, path: str | Path) -> tuple[list[Outcome], list[Problem], Path | None]:
    """Validate everything first; create nothing unless every row passes."""
    path = Path(path)
    projects, problems, _ = check_file(aj, path)
    if problems:
        return [], problems, None
    outcomes = create_all(aj, projects, path)
    return outcomes, [], write_results(path, outcomes)
